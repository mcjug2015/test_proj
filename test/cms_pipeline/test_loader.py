import os
from unittest import mock

from openpyxl.reader.excel import load_workbook
from spark_sql_migrations import custom_logging

from src.cms_pipeline.loader import (
    get_decimal_places,
    get_display_value,
    get_non_empty_cells,
    get_sheet_info_dict,
    get_workbook_sheet_info_dict,
    insert_kvp_rows,
    is_only_text_cell,
    load_zip_workbook,
    parse_sheet,
)

logger = custom_logging.setup_logging().getLogger(__name__)
RES_DIR = os.path.join(os.path.dirname(__file__), "res")


@mock.patch("src.cms_pipeline.loader.convert_to_key", return_value="test converted key")
def test_insert_kvp_rows_success(convert_to_key, migrated_spark, request):
    spark = migrated_spark[0]
    schema = migrated_spark[1]
    logger.info(f"TEST: {request.node.name}; will be using schema {schema};")
    insert_kvp_rows(
        spark,
        cat="spark_catalog",
        schema=schema,
        load_id="test_load_id",
        zip_name="test zip name",
        unzipped_name="test unzipped",
        sheet_name="test sheet name",
        sheet_index=0,
        data_rows=[
            {"test_key": {"value": "test_val", "prev_only_text_cell": "test heading"}}
        ],
    )
    sql_result = spark.sql(f"select * from spark_catalog.{schema}.open_cms_data_kvp;")
    results = [x.asDict() for x in sql_result.toLocalIterator()]
    assert len(results) == 1
    assert results[0]["load_id"] == "test_load_id"
    assert results[0]["table_key"] == "test_key"
    assert results[0]["table_key_simple"] == "test converted key"
    assert results[0]["table_row_index"] == 0
    assert results[0]["table_val"] == "test_val"
    assert results[0]["heading"] == "test heading"
    convert_to_key.assert_called_once()


def test_insert_kvp_rows_no_rows():
    result = insert_kvp_rows(
        None,
        cat=None,
        schema=None,
        load_id=None,
        zip_name=None,
        unzipped_name=None,
        sheet_name=None,
        sheet_index=0,
        data_rows=[],
    )
    assert result == 0


@mock.patch("src.cms_pipeline.loader.download_s3_zip", return_value="/i/am/not/real")
@mock.patch(
    "src.cms_pipeline.loader.Unwrapper.unwrap",
)
@mock.patch("src.cms_pipeline.loader.load_cms_workbook", return_value=1)
@mock.patch("src.cms_pipeline.loader.load_workbook")
def test_load_success(load_workbook, load_cms_workbook, unwrap, download_s3_zip):
    parse_sheet.return_value = (3, [{"test_key": "test_val"}])
    spark = mock.MagicMock(name="spark")
    unwrap.return_value.__enter__.return_value = "/fake/xlsx"

    result = load_zip_workbook(spark, "test_cat", "test_schema", "test_fake_s3_uri")

    assert result == 1
    load_workbook.assert_called_once()
    load_cms_workbook.assert_called_once()
    unwrap.assert_called_once()
    download_s3_zip.assert_called_once()


def test_get_non_empty_cells():
    result = get_non_empty_cells((None, "", "   ", 0, "Year", "  Year  "))

    assert result == ["0", "Year", "Year"]


def test_is_only_text_cell_multiple_cells_returns_false():
    result = is_only_text_cell(["FooterNote", "Metric"])

    assert result is False


def test_is_only_text_cell_single_text_cell_returns_true():
    result = is_only_text_cell(["FooterNote"])

    assert result is True


def test_is_only_text_cell_single_non_text_cell_returns_false():
    result = is_only_text_cell(["12345"])

    assert result is False


def test_get_decimal_places_none_returns_none():
    result = get_decimal_places(None)

    assert result is None


def test_get_decimal_places_integer_format_returns_zero():
    result = get_decimal_places("#,##0")

    assert result == 0


def test_get_decimal_places_two_decimal_format_returns_two():
    result = get_decimal_places("#,##0.00")

    assert result == 2


def test_get_decimal_places_date_format_returns_none():
    result = get_decimal_places("yyyy-mm-dd")

    assert result is None


@mock.patch("src.cms_pipeline.loader.get_decimal_places")
def test_get_display_value_non_float_returns_value_unchanged(get_decimal_places):
    cell = mock.MagicMock(value="Alabama", number_format="General")

    result = get_display_value(cell)

    assert result == "Alabama"
    get_decimal_places.assert_not_called()


@mock.patch("src.cms_pipeline.loader.get_decimal_places", return_value=None)
def test_get_display_value_float_with_no_decimals_info_returns_value_unchanged(
    get_decimal_places,
):
    cell = mock.MagicMock(value=21324800.41667978, number_format="General")

    result = get_display_value(cell)

    assert result == 21324800.41667978
    get_decimal_places.assert_called_once_with("General")


@mock.patch("src.cms_pipeline.loader.get_decimal_places", return_value=0)
def test_get_display_value_float_with_zero_decimals_returns_rounded_int(
    get_decimal_places,
):
    cell = mock.MagicMock(value=21324800.41667978, number_format="#,##0")

    result = get_display_value(cell)

    assert result == 21324800
    assert isinstance(result, int)
    get_decimal_places.assert_called_once_with("#,##0")


@mock.patch("src.cms_pipeline.loader.get_decimal_places", return_value=2)
def test_get_display_value_float_with_nonzero_decimals_returns_rounded_float(
    get_decimal_places,
):
    cell = mock.MagicMock(value=21324800.4166, number_format="#,##0.00")

    result = get_display_value(cell)

    assert result == 21324800.42
    assert isinstance(result, float)
    get_decimal_places.assert_called_once_with("#,##0.00")


def test_get_sheet_info_dict_returns_name_to_description_mapping():
    workbook = load_workbook(os.path.join(RES_DIR, "get_sheet_info_dict_sample.xlsx"))
    result = get_sheet_info_dict(workbook["Table of Contents"])

    # title row (one populated cell), the "Table Name" header row, and the row
    # with only a description (no table name) must all be skipped; whitespace
    # is stripped and extra trailing columns are ignored.
    assert result == {
        "SHEET_A": "Description A",
        "SHEET_B": "Description B",
        "SHEET_C": "Description C",
    }


def test_get_workbook_sheet_info_dict_no_toc():
    workbook = load_workbook(os.path.join(RES_DIR, "parse_sheet_sample.xlsx"))
    result = get_workbook_sheet_info_dict(workbook)

    assert result == {
        "TestSheet": "",
    }


@mock.patch(
    "src.cms_pipeline.loader.get_sheet_info_dict", return_value="testing testing"
)
def test_get_workbook_sheet_info_dict_toc(get_sheet_info_dict):
    workbook = load_workbook(os.path.join(RES_DIR, "get_sheet_info_dict_sample.xlsx"))
    result = get_workbook_sheet_info_dict(workbook)

    assert result == "testing testing"
    get_sheet_info_dict.assert_called_once()


def test_parse_sheet_returns_data_rows():
    workbook = load_workbook(os.path.join(RES_DIR, "parse_sheet_sample.xlsx"))
    data_rows = parse_sheet(workbook["TestSheet"])
    assert data_rows == [
        {
            "Year": {"value": "2023", "prev_only_text_cell": ""},
            "Metric": {"value": "TotalEnroll", "prev_only_text_cell": ""},
            "Value": {"value": "100", "prev_only_text_cell": ""},
        },
        {
            "Year": {"value": "2024", "prev_only_text_cell": ""},
            "Metric": {"value": "TotalEnroll", "prev_only_text_cell": ""},
            "Value": {"value": "200", "prev_only_text_cell": ""},
        },
    ]


def test_parse_sheet_returns_no_rows():
    workbook = load_workbook(os.path.join(RES_DIR, "parse_sheet_no_rows_sample.xlsx"))
    data_rows = parse_sheet(workbook["TestSheet"])

    assert data_rows == []
